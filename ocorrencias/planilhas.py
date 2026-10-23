"""Importação de produtos (.xlsx/.csv) e exportações em Excel (openpyxl)."""

import csv
import io
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .formato import formatar_data
from .labels import (
    CATEGORIAS_COM_SINAL_VOLUME,
    CATEGORIAS_DISPENSAM_DADOS_PRODUTO,
    ITEM_CATEGORIA_LABEL,
    PRIORIDADE_LABEL,
    STATUS_LABEL,
    TIPO_LABEL,
    formatar_referencia,
    rotulo_item,
)

# Alinhado ao limite de itens por ocorrência (MAX_ITENS na criação).
MAX_LINHAS = 100

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------- importação


def _sem_acento(s: str) -> str:
    s = unicodedata.normalize("NFD", s.lower().strip())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def normalizar_categoria(v: str) -> str | None:
    s = _sem_acento(v)
    if s.startswith("avaria"):
        return "avaria"
    if "troca" in s:
        return "troca_volume"
    if s.startswith("falta"):
        return "falta"
    if s.startswith("sobra"):
        return "sobra"
    # "pré-vencido" antes de "vencido" (o segundo é substring do primeiro).
    if "pre" in s and "venc" in s:
        return "pre_vencido"
    if "venc" in s:
        return "vencido"
    return None


def normalizar_sinal(v: str) -> str | None:
    """Coluna "sinal" da planilha (troca de volume): falta ou sobra."""
    s = _sem_acento(v)
    if s.startswith("falt") or s in ("-", "neg"):
        return "falta"
    if s.startswith("sobr") or s in ("+", "pos"):
        return "sobra"
    return None


def _achar_coluna(headers: list[str], candidatos: list[str]) -> int:
    norm = [_sem_acento(h) for h in headers]
    for c in candidatos:
        if c in norm:
            return norm.index(c)
    return -1


def _ler_csv(texto: str) -> list[list[str]]:
    linhas = [linha for linha in texto.splitlines() if linha.strip()]
    if not linhas:
        return []
    # Detecta o separador pela primeira linha.
    sep = ";" if linhas[0].count(";") > linhas[0].count(",") else ","
    return [[c.strip() for c in row] for row in csv.reader(linhas, delimiter=sep)]


def _ler_xlsx(conteudo: bytes) -> list[list[str]]:
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.worksheets[0] if wb.worksheets else None
    if ws is None:
        return []
    saida = []
    for row in ws.iter_rows(values_only=True):
        valores = ["" if c is None else str(c).strip() for c in row]
        # Inteiros lidos como float ("10.0") voltam a ser "10".
        valores = [v[:-2] if re.fullmatch(r"\d+\.0", v) else v for v in valores]
        if any(valores):
            saida.append(valores)
    return saida


def importar_produtos(arquivo) -> dict:
    vazio = {"itens": [], "total": 0, "ignoradas": 0}
    if not arquivo or arquivo.size == 0:
        return {**vazio, "erro": "Arquivo inválido."}

    try:
        conteudo = arquivo.read()
        if arquivo.name.lower().endswith(".csv"):
            try:
                texto = conteudo.decode("utf-8-sig")
            except UnicodeDecodeError:
                texto = conteudo.decode("latin-1")
            linhas = _ler_csv(texto)
        else:
            linhas = _ler_xlsx(conteudo)
    except Exception:  # noqa: BLE001
        return {**vazio, "erro": "Não foi possível ler o arquivo. Use o modelo (.xlsx ou .csv)."}

    if len(linhas) < 2:
        return {**vazio, "erro": "Planilha vazia ou sem linhas de produto."}

    headers = linhas[0]
    i_cod = _achar_coluna(headers, ["codigo", "cod", "sku"])
    i_desc = _achar_coluna(headers, ["nome", "descricao", "produto", "descricao do produto"])
    i_cat = _achar_coluna(headers, ["categoria", "tipo"])
    i_qtd = _achar_coluna(headers, ["quantidade", "qtd", "qtde", "quant"])
    # Opcional: só a troca de volume usa (spec 004).
    i_sinal = _achar_coluna(headers, ["sinal", "falta/sobra", "falta ou sobra", "volume"])

    if min(i_cod, i_desc, i_cat, i_qtd) < 0:
        return {**vazio, "erro": "Colunas esperadas: código, nome, categoria, quantidade."}

    # Não trunca em silêncio: se passar do limite, avisa o cliente.
    if len(linhas) - 1 > MAX_LINHAS:
        return {
            **vazio,
            "erro": f"A planilha tem mais de {MAX_LINHAS} produtos. Divida em ocorrências menores.",
        }

    def celula(linha, i):
        return linha[i].strip() if i < len(linha) else ""

    itens, ignoradas = [], 0
    for linha in linhas[1:]:
        codigo = celula(linha, i_cod)
        descricao = celula(linha, i_desc)
        texto_categoria = celula(linha, i_cat)
        categoria = normalizar_categoria(texto_categoria)
        digitos = re.sub(r"\D", "", celula(linha, i_qtd))
        quantidade = int(digitos) if digitos else 0
        sinal_volume = None
        if categoria in CATEGORIAS_COM_SINAL_VOLUME:
            # Coluna "sinal" ou, na falta dela, o texto da categoria ("Troca de volume - falta").
            texto_sinal = celula(linha, i_sinal) if i_sinal >= 0 else ""
            resto_categoria = _sem_acento(texto_categoria).replace("troca de volume", "").replace("troca", "")
            sinal_volume = normalizar_sinal(texto_sinal) or normalizar_sinal(resto_categoria.strip(" -:/()"))
            if sinal_volume is None:
                ignoradas += 1
                continue
        if categoria in CATEGORIAS_DISPENSAM_DADOS_PRODUTO:
            # Mesma regra do formulário (spec 009): código, descrição e quantidade opcionais.
            descricao = descricao or ITEM_CATEGORIA_LABEL[categoria]
            quantidade = quantidade or 1
            codigo_obrigatorio = False
        else:
            codigo_obrigatorio = True
        if not categoria or (codigo_obrigatorio and not codigo) or not descricao or quantidade <= 0:
            ignoradas += 1
            continue
        itens.append(
            {
                "categoria": categoria,
                "codigo": codigo,
                "descricao": descricao,
                "quantidade": quantidade,
                "sinal_volume": sinal_volume,
            }
        )

    return {"itens": itens, "total": len(itens), "ignoradas": ignoradas}


# ---------------------------------------------------------------- exportação


def _cabecalho(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF18181B")


def _larguras(ws, larguras: list[int]):
    for i, w in enumerate(larguras):
        ws.column_dimensions[chr(ord("A") + i)].width = w


def _bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def planilha_ocorrencia(oc) -> bytes:
    """Resumo + produtos de uma ocorrência."""
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    resumo.append(["Campo", "Valor"])
    _cabecalho(resumo)
    _larguras(resumo, [20, 50])
    cliente = oc.cliente
    for campo, valor in [
        ("Protocolo", oc.protocolo),
        ("Tipo", TIPO_LABEL.get(oc.tipo, oc.tipo)),
        ("Referência", formatar_referencia(oc.referencia_tipo, oc.referencia) or "—"),
        ("Status", STATUS_LABEL.get(oc.status, oc.status)),
        ("Cliente", cliente.nome if cliente else "—"),
        ("Empresa", cliente.empresa.nome if cliente and cliente.empresa_id else "—"),
        ("Aberta em", formatar_data(oc.criado_em)),
        ("Descrição", oc.descricao or "—"),
    ]:
        resumo.append([campo, valor])

    itens = list(oc.itens.all())
    if itens:
        ws = wb.create_sheet("Produtos")
        ws.append(["Categoria", "Código", "Descrição", "Quantidade"])
        _cabecalho(ws)
        _larguras(ws, [18, 16, 40, 12])
        for it in itens:
            ws.append([rotulo_item(it.categoria, it.sinal_volume), it.codigo, it.descricao, it.quantidade])
        ws.auto_filter.ref = f"A1:D{len(itens) + 1}"

    return _bytes(wb)


def planilha_ocorrencias(ocorrencias) -> bytes:
    """Lista filtrada do painel da operação."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ocorrências"
    ws.append(["Protocolo", "Cliente", "Empresa", "Categoria", "Status", "Prioridade", "Referência", "Aberta em"])
    _cabecalho(ws)
    _larguras(ws, [16, 24, 24, 18, 14, 12, 20, 20])
    n = 0
    for oc in ocorrencias:
        n += 1
        ws.append(
            [
                oc.protocolo,
                oc.cliente.nome,
                oc.cliente.empresa.nome if oc.cliente.empresa_id else "",
                TIPO_LABEL.get(oc.tipo, oc.tipo),
                STATUS_LABEL.get(oc.status, oc.status),
                PRIORIDADE_LABEL.get(oc.prioridade, oc.prioridade),
                formatar_referencia(oc.referencia_tipo, oc.referencia),
                formatar_data(oc.criado_em),
            ]
        )
    ws.auto_filter.ref = f"A1:H{n + 1}"
    return _bytes(wb)
