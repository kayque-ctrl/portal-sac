"""Operação: fila, atendimento, ciclo de vida, notificações e painel."""

from django.utils import timezone

from ocorrencias.models import Notificacao, Ocorrencia

from ..builders import foto

# ---------------------------------------------------------------- fila


def corpo(resposta) -> str:
    return resposta.content.decode()


def test_fila_filtra_por_status_tipo_e_empresa(ocorrencia, abrir, estranho, operador, empresa, entrar):
    outra = abrir(por=estranho)
    http = entrar(operador)

    assert ocorrencia.protocolo in corpo(http.get(f"/operacao/?tipo=produtos&empresa={empresa.pk}"))
    assert outra.protocolo not in corpo(http.get(f"/operacao/?empresa={empresa.pk}"))
    assert outra.protocolo in corpo(http.get("/operacao/?status=aberta&tipo=atraso"))


def test_fila_ignora_empresa_invalida(ocorrencia, operador, entrar):
    assert ocorrencia.protocolo in corpo(entrar(operador).get("/operacao/?empresa=nao-e-uuid"))


def test_fila_busca_por_nome_do_cliente_e_pela_referencia(ocorrencia, operador, entrar):
    http = entrar(operador)
    assert ocorrencia.protocolo in corpo(http.get("/operacao/?q=Ana"))
    assert ocorrencia.protocolo in corpo(http.get("/operacao/?q=12345"))
    assert ocorrencia.protocolo not in corpo(http.get("/operacao/?q=inexistente"))


def test_abas_arquivadas_concluidas_e_lixeira_separam_as_ocorrencias(abrir, operador, admin, entrar):
    arquivada, concluida, lixo = abrir(), abrir(), abrir()
    Ocorrencia.objects.filter(pk=arquivada.pk).update(arquivada_em=timezone.now())
    Ocorrencia.objects.filter(pk=concluida.pk).update(status="recusada")
    Ocorrencia.objects.filter(pk=lixo.pk).update(excluida_em=timezone.now())
    http = entrar(admin)

    assert arquivada.protocolo in corpo(http.get("/operacao/?situacao=arquivadas"))
    assert concluida.protocolo in corpo(http.get("/operacao/?situacao=concluidas"))
    assert lixo.protocolo in corpo(http.get("/operacao/?situacao=lixeira"))
    assert lixo.protocolo not in corpo(http.get("/operacao/"))


def test_operador_nao_acessa_a_aba_lixeira_nem_por_url(abrir, operador, entrar):
    lixo = abrir()
    Ocorrencia.objects.filter(pk=lixo.pk).update(excluida_em=timezone.now())

    resposta = entrar(operador).get("/operacao/?situacao=lixeira")

    assert lixo.protocolo not in corpo(resposta)


def test_admin_esvazia_a_lixeira(abrir, admin, entrar):
    lixo, ativa = abrir(), abrir()
    Ocorrencia.objects.filter(pk=lixo.pk).update(excluida_em=timezone.now())

    r = entrar(admin).post("/operacao/lixeira/esvaziar")

    assert r.url == "/operacao/?situacao=lixeira"
    assert not Ocorrencia.objects.filter(pk=lixo.pk).exists()
    assert Ocorrencia.objects.filter(pk=ativa.pk).exists()


# ---------------------------------------------------------------- atendimento


def test_atualizar_com_sucesso_mostra_confirmacao(ocorrencia, operador, entrar):
    http = entrar(operador)
    r = http.post(
        f"/operacao/{ocorrencia.pk}/atualizar",
        {"status": "em_analise", "prioridade": "alta", "atribuido_a": str(operador.pk)},
        follow=True,
    )
    assert "Alterações salvas." in corpo(r)
    ocorrencia.refresh_from_db()
    assert (ocorrencia.status, ocorrencia.prioridade, ocorrencia.atribuido_a) == ("em_analise", "alta", operador)


def test_mensagem_da_operacao_publica_e_interna(ocorrencia, operador, entrar):
    http = entrar(operador)

    http.post(f"/operacao/{ocorrencia.pk}/mensagem", {"conteudo": "Resposta ao cliente"})
    http.post(f"/operacao/{ocorrencia.pk}/mensagem", {"conteudo": "Anotação", "interna": "on"})

    assert ocorrencia.mensagens.filter(conteudo="Resposta ao cliente", interna=False).exists()
    assert ocorrencia.mensagens.filter(conteudo="Anotação", interna=True).exists()


def test_mensagem_vazia_da_operacao_mostra_o_erro_e_guarda_o_marcador_interno(ocorrencia, operador, entrar):
    r = entrar(operador).post(f"/operacao/{ocorrencia.pk}/mensagem", {"conteudo": "", "interna": "on"})
    assert "Escreva uma mensagem ou anexe um arquivo." in corpo(r)


def test_mensagem_com_anexo_que_falha_avisa(ocorrencia, operador, entrar):
    from unittest import mock

    from ocorrencias.models import Anexo

    with mock.patch.object(Anexo.objects, "create", side_effect=OSError("disco")):
        r = entrar(operador).post(
            f"/operacao/{ocorrencia.pk}/mensagem", {"conteudo": "com anexo", "anexos": [foto("x.png")]}, follow=True
        )

    assert "não foi possível anexar" in corpo(r)


def test_coleta_com_email_invalido_mostra_o_erro(ocorrencia, operador, entrar):
    r = entrar(operador).post(f"/operacao/{ocorrencia.pk}/coleta", {"email_transportadora": "invalido"})
    assert "Informe um e-mail válido da transportadora." in corpo(r)


def test_coleta_registrada_mostra_confirmacao(ocorrencia, operador, entrar):
    r = entrar(operador).post(
        f"/operacao/{ocorrencia.pk}/coleta", {"email_transportadora": "transp@exemplo.com"}, follow=True
    )
    assert "Solicitação registrada." in corpo(r)


# ---------------------------------------------------------------- ciclo de vida


def test_arquivar_e_desarquivar(ocorrencia, operador, entrar):
    http = entrar(operador)

    http.post(f"/operacao/{ocorrencia.pk}/ciclo/arquivar")
    ocorrencia.refresh_from_db()
    assert ocorrencia.arquivada_em is not None

    http.post(f"/operacao/{ocorrencia.pk}/ciclo/desarquivar")
    ocorrencia.refresh_from_db()
    assert ocorrencia.arquivada_em is None


def test_restaurar_da_lixeira(ocorrencia, admin, entrar):
    http = entrar(admin)
    http.post(f"/operacao/{ocorrencia.pk}/ciclo/excluir")

    http.post(f"/operacao/{ocorrencia.pk}/ciclo/restaurar")

    ocorrencia.refresh_from_db()
    assert ocorrencia.excluida_em is None


def test_operacao_reabre_ocorrencia_encerrada(ocorrencia, operador, entrar):
    ocorrencia.status = "recusada"
    ocorrencia.motivo_recusa = "x"
    ocorrencia.save(por=operador)

    entrar(operador).post(f"/operacao/{ocorrencia.pk}/ciclo/reabrir")

    ocorrencia.refresh_from_db()
    assert ocorrencia.status == "em_analise"


def test_acao_de_ciclo_desconhecida_e_404(ocorrencia, operador, entrar):
    assert entrar(operador).post(f"/operacao/{ocorrencia.pk}/ciclo/explodir").status_code == 404


def test_ocorrencia_inexistente_e_404(operador, entrar):
    import uuid

    assert entrar(operador).get(f"/operacao/{uuid.uuid4()}").status_code == 404


def test_exportar_ocorrencia_pela_operacao(ocorrencia, operador, entrar):
    r = entrar(operador).get(f"/operacao/{ocorrencia.pk}/export")
    assert r.content.startswith(b"PK")
    assert ocorrencia.protocolo in r["Content-Disposition"]


# ---------------------------------------------------------------- notificações


def test_pagina_de_notificacoes_marca_como_lidas(ocorrencia, operador, entrar):
    assert operador.notificacoes.filter(lida=False).exists()

    entrar(operador).get("/operacao/notificacoes")

    assert not operador.notificacoes.filter(lida=False).exists()


def test_sino_lista_recentes_com_link_para_a_area_certa(ocorrencia, operador, cliente, entrar):
    dados = entrar(operador).get("/notificacoes/recentes").json()
    assert dados["naoLidas"] >= 1
    assert dados["itens"][0]["href"] == f"/operacao/{ocorrencia.pk}"

    Notificacao.objects.create(destinatario=cliente, ocorrencia=ocorrencia, mensagem="Sua ocorrência está em tratativa")
    dados = entrar(cliente).get("/notificacoes/recentes").json()
    assert dados["itens"][0]["href"] == f"/portal/{ocorrencia.pk}"


def test_notificacao_sem_ocorrencia_nao_tem_link(cliente, entrar):
    Notificacao.objects.create(destinatario=cliente, mensagem="Aviso geral")
    assert entrar(cliente).get("/notificacoes/recentes").json()["itens"][0]["href"] is None


def test_marcar_lidas_so_marca_as_informadas_e_do_proprio_usuario(ocorrencia, cliente, colega, entrar):
    minha = Notificacao.objects.create(destinatario=cliente, mensagem="a")
    outra_minha = Notificacao.objects.create(destinatario=cliente, mensagem="b")
    do_colega = Notificacao.objects.create(destinatario=colega, mensagem="c")

    entrar(cliente).post("/notificacoes/marcar-lidas", {"ids": [str(minha.pk), str(do_colega.pk), "lixo"]})

    for n in (minha, outra_minha, do_colega):
        n.refresh_from_db()
    assert (minha.lida, outra_minha.lida, do_colega.lida) == (True, False, False)


def test_marcar_lidas_sem_ids_nao_faz_nada(cliente, entrar):
    n = Notificacao.objects.create(destinatario=cliente, mensagem="a")
    assert entrar(cliente).post("/notificacoes/marcar-lidas").json() == {"ok": True}
    n.refresh_from_db()
    assert not n.lida


# ---------------------------------------------------------------- painel


def test_painel_aceita_filtros_e_ignora_valores_invalidos(ocorrencia, operador, entrar):
    http = entrar(operador)
    assert http.get("/operacao/painel?status=aberta&tipo=produtos&de=2020-01-01&ate=2099-12-31").status_code == 200
    assert http.get("/operacao/painel?status=xx&tipo=yy&de=ontem&ate=amanha&empresa=zz").status_code == 200


def test_exportacao_do_painel_respeita_os_filtros(ocorrencia, abrir, estranho, operador, empresa, entrar):
    from io import BytesIO

    from openpyxl import load_workbook

    abrir(por=estranho)
    r = entrar(operador).get(
        f"/operacao/painel/export?empresa={empresa.pk}&de=2020-01-01&ate=2099-12-31&tipo=produtos&status=em_analise"
    )

    linhas = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
    assert [linha[0] for linha in linhas[1:]] == [ocorrencia.protocolo]
    assert linhas[1][2] == "ACME"
