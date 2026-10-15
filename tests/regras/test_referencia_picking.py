"""Spec 016: número de picking como terceira opção de referência."""

import pytest

from ocorrencias import labels, servicos

from ..builders import dados_atraso, dados_produtos, dados_reclamacao


@pytest.mark.parametrize("tipo", ["nf", "pedido", "picking"])
def test_toda_referencia_do_enum_tem_rotulo_e_exemplo(tipo):
    form = labels.REFERENCIA_FORM[tipo]
    assert form["campo"] and form["exemplo"].startswith("Ex.:") and form["curto"]


def test_enum_e_formulario_tem_os_mesmos_tipos():
    assert set(labels.REFERENCIA_FORM) == set(labels.ReferenciaTipo.values) == {"nf", "pedido", "picking"}


@pytest.mark.parametrize(
    ("tipo", "esperado"),
    [("nf", "NF 123"), ("pedido", "Pedido 123"), ("picking", "Picking 123"), (None, "123")],
)
def test_formatar_referencia(tipo, esperado):
    assert labels.formatar_referencia(tipo, "123") == esperado


@pytest.mark.parametrize("dados", [dados_atraso, dados_produtos], ids=["atraso", "produtos"])
def test_abre_ocorrencia_com_picking(cliente, dados):
    r = servicos.criar_ocorrencia(cliente, dados(referencia_tipo="picking", referencia="004512"), [])

    assert r.ok, r.erro
    oc = r.dados["ocorrencia"]
    assert (oc.referencia_tipo, oc.referencia, oc.referencia_formatada) == ("picking", "004512", "Picking 004512")


def test_picking_sem_numero_pede_o_numero_do_picking(cliente):
    r = servicos.criar_ocorrencia(cliente, dados_atraso(referencia_tipo="picking", referencia=""), [])
    assert r.erro == "Informe o número do picking."


def test_tipo_de_referencia_desconhecido_continua_recusado(cliente):
    r = servicos.criar_ocorrencia(cliente, dados_atraso(referencia_tipo="romaneio"), [])
    assert r.erro == "Escolha a referência: Nota Fiscal, Pedido ou Picking."


def test_picking_opcional_na_reclamacao(cliente):
    oc = servicos.criar_ocorrencia(cliente, dados_reclamacao(referencia_tipo="picking", referencia="77"), []).dados["ocorrencia"]
    assert oc.referencia_formatada == "Picking 77"


def test_duplicidade_considera_o_picking(abrir, colega):
    abrir(referencia_tipo="picking", referencia="004512")
    assert servicos.verificar_duplicata(colega, "004512")["abertaPor"] == "Ana Cliente"


def test_exportacao_mostra_o_picking(cliente):
    from io import BytesIO

    from openpyxl import load_workbook

    from ocorrencias.planilhas import planilha_ocorrencia

    oc = servicos.criar_ocorrencia(cliente, dados_atraso(referencia_tipo="picking", referencia="004512"), []).dados["ocorrencia"]
    resumo = {linha[0]: linha[1] for linha in load_workbook(BytesIO(planilha_ocorrencia(oc)))["Resumo"].iter_rows(values_only=True)}

    assert resumo["Referência"] == "Picking 004512"
